"""Document text extraction (PDF, DOCX, PPTX, HTML, Markdown/plain text) into heading-annotated text."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED = {"pdf", "docx", "pptx", "html", "htm", "txt", "md", "markdown", "text"}


@dataclass
class ExtractedDocument:
    text: str
    title: str | None = None
    pages: int = 0
    meta: dict = field(default_factory=dict)


class ExtractionError(ValueError):
    pass


def kind_from_name(filename: str | None, content_type: str | None = None) -> str:
    ext = Path(filename or "").suffix.lower().lstrip(".")
    if ext in SUPPORTED:
        return {"htm": "html", "markdown": "md", "text": "txt"}.get(ext, ext)
    ct = (content_type or "").split(";")[0].strip().lower()
    return {
        "application/pdf": "pdf",
        "text/html": "html",
        "application/xhtml+xml": "html",
        "text/markdown": "md",
        "text/plain": "txt",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    }.get(ct, "txt")


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16") if data[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig",):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n[ \t]+", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def extract_pdf(data: bytes) -> ExtractedDocument:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:  # noqa: BLE001
                raise ExtractionError("PDF is password protected") from exc
        pages = []
        for i, page in enumerate(reader.pages, 1):
            body = (page.extract_text() or "").strip()
            # Page markers help citations; de-hyphenate line-wrapped words.
            body = re.sub(r"(\w)-\n(\w)", r"\1\2", body)
            if body:
                pages.append(f"<!-- page {i} -->\n{body}")
        title = None
        try:
            title = (reader.metadata.title if reader.metadata else None) or None
        except Exception:  # noqa: BLE001 - metadata parsing is best-effort
            title = None
        return ExtractedDocument(_clean("\n\n".join(pages)), title, len(reader.pages))
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"Could not read PDF: {exc}") from exc


def extract_docx(data: bytes) -> ExtractedDocument:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"Could not read DOCX: {exc}") from exc
    lines: list[str] = []
    title = (document.core_properties.title or None) if document.core_properties else None
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name if para.style is not None else "") or ""
        m = re.match(r"Heading (\d)", style)
        if style == "Title":
            title = title or text
            lines.append(f"# {text}")
        elif m:
            lines.append(f"{'#' * min(6, int(m.group(1)) + (1 if title else 0))} {text}")
        elif style.startswith("List"):
            lines.append(f"- {text}")
        else:
            lines.append(text)
        lines.append("")
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
        lines.append("")
    return ExtractedDocument(_clean("\n".join(lines)), title)


def extract_pptx(data: bytes) -> ExtractedDocument:
    from pptx import Presentation

    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"Could not read PPTX: {exc}") from exc
    lines: list[str] = []
    title = None
    for i, slide in enumerate(prs.slides, 1):
        slide_title = None
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            slide_title = slide.shapes.title.text_frame.text.strip() or None
        title = title or slide_title
        lines.append(f"## Slide {i}" + (f": {slide_title}" if slide_title else ""))
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if t:
                        lines.append(("  " * p.level) + f"- {t}")
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    lines.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide:
            notes = (
                slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            )
            if notes:
                lines.append(f"Speaker notes: {notes}")
        lines.append("")
    return ExtractedDocument(_clean("\n".join(lines)), title, len(prs.slides))


def extract_html(data: bytes | str) -> ExtractedDocument:
    from bs4 import BeautifulSoup

    html = data if isinstance(data, str) else _decode(data)
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(
        ["script", "style", "noscript", "nav", "footer", "header", "aside", "form", "iframe", "svg"]
    ):
        tag.decompose()
    title = soup.title.get_text(strip=True) if soup.title else None
    root = soup.find("main") or soup.find("article") or soup.body or soup
    lines: list[str] = []
    for el in root.find_all(
        ["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre", "td", "th", "blockquote", "dt", "dd"]
    ):
        if el.find_parent(["li", "td", "th"]) is not None and el.name in ("p",):
            continue
        text = el.get_text(" ", strip=True)
        if not text:
            continue
        if el.name[0] == "h" and el.name[1:].isdigit():
            lines.append(f"{'#' * int(el.name[1])} {text}")
        elif el.name == "li":
            lines.append(f"- {text}")
        else:
            lines.append(text)
        lines.append("")
    if not lines:
        lines = [root.get_text("\n", strip=True)]
    return ExtractedDocument(_clean("\n".join(lines)), title)


def extract_text(data: bytes, kind: str) -> ExtractedDocument:
    """Extract text for ``kind`` in {pdf, docx, pptx, html, md, txt}. Raises ``ExtractionError``."""
    kind = {"htm": "html", "markdown": "md", "text": "txt", "url": "html"}.get(kind, kind)
    if kind == "pdf":
        doc = extract_pdf(data)
    elif kind == "docx":
        doc = extract_docx(data)
    elif kind == "pptx":
        doc = extract_pptx(data)
    elif kind == "html":
        doc = extract_html(data)
    else:
        text = _decode(data)
        title = None
        m = re.search(r"^#\s+(.+)$", text, re.M)
        if m:
            title = m.group(1).strip()
        doc = ExtractedDocument(_clean(text), title)
    if not doc.text.strip():
        raise ExtractionError("No extractable text found (scanned or empty document?)")
    return doc
