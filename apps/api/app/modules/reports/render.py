"""Report renderers: CSV, XLSX (openpyxl) and branded PDF (reportlab)."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import Any
from uuid import UUID

DEFAULT_BRAND = "#4f46e5"
_HEX = re.compile(r"^#[0-9a-fA-F]{6}$")
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


@dataclass
class ReportTable:
    title: str
    columns: list[str]
    rows: list[list[Any]]


@dataclass
class ReportData:
    title: str
    subtitle: str
    organisation_name: str
    brand_color: str = DEFAULT_BRAND
    kpis: list[tuple[str, Any]] = field(default_factory=list)
    tables: list[ReportTable] = field(default_factory=list)
    filters: dict[str, Any] = field(default_factory=dict)
    generated_at: datetime | None = None
    generated_by: str | None = None

    @property
    def row_count(self) -> int:
        return len(self.tables[0].rows) if self.tables else 0


def brand(color: Any) -> str:
    return color if isinstance(color, str) and _HEX.match(color) else DEFAULT_BRAND


def display(value: Any) -> Any:
    """Normalise a cell for output (keeps numbers numeric for XLSX)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, float):
        return round(value, 2)
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, UUID):
        return str(value)
    if isinstance(value, list | tuple):
        return ", ".join(str(v) for v in value)
    return value


def safe_text(value: Any) -> Any:
    """Neutralise spreadsheet formula injection in user-controlled strings."""
    v = display(value)
    if isinstance(v, str) and v.startswith(_FORMULA_PREFIXES):
        return "'" + v
    return v


# ----------------------------------------------------------------------------- CSV


def to_csv(report: ReportData) -> bytes:
    """CSV carries the report's primary (first) table; XLSX/PDF include every section and the KPIs."""
    buf = io.StringIO()
    writer = csv.writer(buf)
    table = (
        report.tables[0]
        if report.tables
        else ReportTable("Summary", ["Metric", "Value"], [list(k) for k in report.kpis])
    )
    writer.writerow(table.columns)
    for row in table.rows:
        writer.writerow([safe_text(v) for v in row])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel opens UTF-8 correctly


# ----------------------------------------------------------------------------- XLSX


def _sheet_title(title: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]\:\*\?\/\\]", " ", title).strip()[:28] or "Sheet"
    name, i = base, 2
    while name.lower() in used:
        name = f"{base[:25]} ({i})"
        i += 1
    used.add(name.lower())
    return name


def to_xlsx(report: ReportData) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    color = brand(report.brand_color).lstrip("#").upper()
    header_fill = PatternFill("solid", start_color=color, end_color=color)
    header_font = Font(bold=True, color="FFFFFF")
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    used = {"summary"}
    ws["A1"] = report.title
    ws["A1"].font = Font(bold=True, size=14, color=color)
    ws["A2"] = report.organisation_name
    ws["A3"] = report.subtitle
    ws["A4"] = f"Generated {display(report.generated_at)}" + (
        f" by {report.generated_by}" if report.generated_by else ""
    )
    r = 6
    for label, value in report.filters.items():
        if value not in (None, ""):
            ws.cell(row=r, column=1, value=f"Filter: {label}")
            ws.cell(row=r, column=2, value=safe_text(value))
            r += 1
    r += 1
    ws.cell(row=r, column=1, value="Metric").font = header_font
    ws.cell(row=r, column=1).fill = header_fill
    ws.cell(row=r, column=2, value="Value").font = header_font
    ws.cell(row=r, column=2).fill = header_fill
    for label, value in report.kpis:
        r += 1
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=safe_text(value))
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 28

    for table in report.tables:
        sheet = wb.create_sheet(_sheet_title(table.title, used))
        sheet.append(table.columns)
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
        widths = [len(str(c)) for c in table.columns]
        for row in table.rows:
            values = [safe_text(v) for v in row]
            sheet.append(values)
            for i, v in enumerate(values):
                widths[i] = max(widths[i], len(str(v)))
        for i, w in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(i)].width = min(max(10, w + 2), 60)
        sheet.freeze_panes = "A2"
        if table.rows:
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(table.columns))}{len(table.rows) + 1}"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ----------------------------------------------------------------------------- PDF

MAX_PDF_ROWS = 2000


def to_pdf(report: ReportData) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    primary = colors.HexColor(brand(report.brand_color))
    page = landscape(A4)
    margin = 14 * mm
    header_h = 22 * mm
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=page,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=header_h + 8 * mm,
        bottomMargin=14 * mm,
        title=report.title,
        author=report.organisation_name,
    )
    width = page[0] - 2 * margin
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=7.5, leading=9)
    head_style = ParagraphStyle("head", parent=cell_style, textColor=colors.white, fontName="Helvetica-Bold")
    h2 = ParagraphStyle(
        "h2", parent=styles["Heading2"], textColor=primary, fontSize=12, spaceBefore=8, spaceAfter=4
    )
    small = ParagraphStyle(
        "small", parent=styles["BodyText"], fontSize=8, textColor=colors.HexColor("#4b5563")
    )

    def esc(v: Any) -> str:
        return str(display(v)).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    def on_page(canvas: Any, _doc: Any) -> None:
        canvas.saveState()
        canvas.setFillColor(primary)
        canvas.rect(0, page[1] - header_h, page[0], header_h, stroke=0, fill=1)
        canvas.setFillColor(colors.white)
        canvas.setFont("Helvetica-Bold", 15)
        canvas.drawString(margin, page[1] - 11 * mm, report.organisation_name[:80])
        canvas.setFont("Helvetica", 10)
        canvas.drawString(margin, page[1] - 17 * mm, report.title)
        canvas.drawRightString(
            page[0] - margin, page[1] - 11 * mm, f"Generated {display(report.generated_at)}"
        )
        canvas.setFillColor(colors.HexColor("#6b7280"))
        canvas.setFont("Helvetica", 8)
        canvas.drawString(margin, 8 * mm, "Confidential — generated by Praxis")
        canvas.drawRightString(page[0] - margin, 8 * mm, f"Page {canvas.getPageNumber()}")
        canvas.restoreState()

    story: list[Any] = [Paragraph(esc(report.subtitle), small)]
    active_filters = [f"{k}: {esc(v)}" for k, v in report.filters.items() if v not in (None, "")]
    if active_filters:
        story.append(Paragraph("Filters — " + "; ".join(active_filters), small))
    story.append(Spacer(1, 4 * mm))

    if report.kpis:
        per_row = 4
        cells = []
        for label, value in report.kpis:
            cells.append(
                Paragraph(
                    f'<font size="7" color="#6b7280">{esc(label).upper()}</font><br/>'
                    f'<font size="14"><b>{esc(value)}</b></font>',
                    cell_style,
                )
            )
        while len(cells) % per_row:
            cells.append("")
        grid = [cells[i : i + per_row] for i in range(0, len(cells), per_row)]
        kpi_table = Table(grid, colWidths=[width / per_row] * per_row, rowHeights=None)
        kpi_table.setStyle(
            TableStyle(
                [
                    ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#e5e7eb")),
                    ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e5e7eb")),
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f9fafb")),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                    ("LINEABOVE", (0, 0), (-1, 0), 2, primary),
                ]
            )
        )
        story += [kpi_table, Spacer(1, 4 * mm)]

    for table in report.tables:
        story.append(Paragraph(esc(table.title), h2))
        if not table.rows:
            story.append(Paragraph("No data for the selected filters.", small))
            continue
        rows = table.rows[:MAX_PDF_ROWS]
        data = [[Paragraph(esc(c), head_style) for c in table.columns]]
        data += [[Paragraph(esc(v), cell_style) for v in row] for row in rows]
        t = Table(data, colWidths=[width / len(table.columns)] * len(table.columns), repeatRows=1)
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), primary),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        story.append(t)
        if len(table.rows) > MAX_PDF_ROWS:
            story.append(
                Paragraph(
                    f"Showing the first {MAX_PDF_ROWS} of {len(table.rows)} rows — export to "
                    "Excel or CSV for the full data set.",
                    small,
                )
            )
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


RENDERERS = {
    "csv": (to_csv, "text/csv; charset=utf-8"),
    "xlsx": (to_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "pdf": (to_pdf, "application/pdf"),
}
