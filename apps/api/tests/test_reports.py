"""Reports: every type × format renders and parses back, permissions (incl. learner self-service), filters,
tenant isolation, formula-injection safety and the report.generated audit entry."""

from __future__ import annotations

import csv
import io

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select

from app.core.db import get_sessionmaker
from app.models.platform import AuditLog
from app.modules.reports.render import ReportData, ReportTable, to_csv, to_xlsx
from tests.factories import add_member, signup
from tests.ops_seed import seed_scenario

API = "/api/v1/reports"
TYPES = ["learner_progress", "course", "compliance", "assignment", "simulation", "competency", "organisation"]


@pytest.fixture
async def scenario(client):
    admin = await signup(client, org_name="Reports Org")
    sc = await seed_scenario(admin.org_id, admin.user_id)
    return admin, sc


def _csv_rows(content: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))


@pytest.mark.parametrize("fmt", ["pdf", "xlsx", "csv"])
async def test_all_types_render(client, scenario, fmt):
    admin, sc = scenario
    for report_type in TYPES:
        body = {"format": fmt, "course_id": str(sc.course.id)}
        r = await client.post(f"{API}/{report_type}", json=body, headers=admin.headers)
        assert r.status_code == 200, (report_type, r.text[:300])
        disp = r.headers["content-disposition"]
        assert disp.startswith("attachment;") and disp.endswith(f'.{fmt}"'), disp
        if fmt == "pdf":
            assert r.content.startswith(b"%PDF")
            text = "".join(p.extract_text() for p in PdfReader(io.BytesIO(r.content)).pages)
            assert "Reports Org" in text
        elif fmt == "xlsx":
            wb = load_workbook(io.BytesIO(r.content))
            assert wb.sheetnames[0] == "Summary" and len(wb.sheetnames) >= 2
            sheet = wb[wb.sheetnames[1]]
            assert sheet.freeze_panes == "A2" and sheet["A1"].font.bold
        else:
            rows = _csv_rows(r.content)
            assert len(rows) >= 1 and rows[0]


async def test_learner_progress_content_and_filters(client, scenario):
    admin, sc = scenario
    r = await client.post(f"{API}/learner_progress", json={"format": "csv"}, headers=admin.headers)
    rows = _csv_rows(r.content)
    header, data = rows[0], rows[1:]
    assert header[:5] == ["Learner", "Email", "Department", "Course", "Status"]
    assert sorted(row[0] for row in data) == ["Learner 1", "Learner 2", "Learner 3"]
    assert r.headers["x-report-rows"] == "3"
    # department filter
    r = await client.post(
        f"{API}/learner_progress",
        json={"format": "csv", "department_id": str(sc.department_id)},
        headers=admin.headers,
    )
    assert sorted(row[0] for row in _csv_rows(r.content)[1:]) == ["Learner 1", "Learner 2"]
    # format via query string; xlsx summary has KPIs
    r = await client.post(f"{API}/learner_progress?format=xlsx", json={}, headers=admin.headers)
    wb = load_workbook(io.BytesIO(r.content))
    values = {row[0]: row[1] for row in wb["Summary"].iter_rows(values_only=True) if row[0]}
    assert values["Enrolments"] == 3 and values["Completed"] == 1 and values["Overdue"] == 1


async def test_compliance_and_competency_numbers(client, scenario):
    admin, sc = scenario
    r = await client.post(f"{API}/compliance", json={"format": "csv"}, headers=admin.headers)
    rows = _csv_rows(r.content)
    by_learner = {row[0]: row for row in rows[1:]}
    assert set(by_learner) == {"Learner 1", "Learner 2"}
    idx = rows[0].index("Compliance")
    assert by_learner["Learner 1"][idx] == "On time" and by_learner["Learner 2"][idx] == "Overdue"
    assert by_learner["Learner 2"][rows[0].index("Days overdue")] == "1"

    r = await client.post(f"{API}/competency", json={"format": "xlsx"}, headers=admin.headers)
    wb = load_workbook(io.BytesIO(r.content))
    gap_rows = list(wb[wb.sheetnames[1]].iter_rows(values_only=True))
    assert gap_rows[1][0] == "Network Admin" and gap_rows[1][6] == 3  # below


async def test_course_report_requires_course(client, scenario):
    admin, _ = scenario
    r = await client.post(f"{API}/course", json={"format": "pdf"}, headers=admin.headers)
    assert r.status_code == 400
    assert (await client.post(f"{API}/bogus", json={}, headers=admin.headers)).status_code == 422
    assert (
        await client.post(f"{API}/organisation", json={"format": "docx"}, headers=admin.headers)
    ).status_code == 422


async def test_permissions_and_self_service(client, scenario):
    admin, sc = scenario
    l1, l2, _ = sc.learners
    assert (await client.post(f"{API}/organisation", json={}, headers=l1.headers)).status_code == 403
    assert (
        await client.post(f"{API}/learner_progress", json={"user_id": str(l2.user_id)}, headers=l1.headers)
    ).status_code == 403
    r = await client.post(f"{API}/learner_progress", json={"format": "csv"}, headers=l1.headers)
    assert r.status_code == 200
    assert [row[0] for row in _csv_rows(r.content)[1:]] == ["Learner 1"]
    types = [t["type"] for t in (await client.get(f"{API}/types", headers=l1.headers)).json()]
    assert types == ["learner_progress"]
    auditor = await add_member(admin.org_id, "auditor")
    assert (
        await client.post(f"{API}/compliance", json={"format": "csv"}, headers=auditor.headers)
    ).status_code == 200
    assert len((await client.get(f"{API}/types", headers=auditor.headers)).json()) == 7

    other = await signup(client, org_name="Other Reports Org")
    r = await client.post(f"{API}/course", json={"course_id": str(sc.course.id)}, headers=other.headers)
    assert r.status_code == 404
    r = await client.post(f"{API}/learner_progress", json={"format": "csv"}, headers=other.headers)
    assert _csv_rows(r.content)[1:] == []

    async with get_sessionmaker()() as s:
        s.info["rls_bypass"] = True
        logs = (
            (
                await s.execute(
                    select(AuditLog).where(
                        AuditLog.organisation_id == admin.org_id, AuditLog.action == "report.generated"
                    )
                )
            )
            .scalars()
            .all()
        )
    assert logs and all("report_type" in log.meta for log in logs)


def test_formula_injection_is_neutralised():
    report = ReportData(
        title="T",
        subtitle="S",
        organisation_name="O",
        kpis=[("k", 1)],
        tables=[ReportTable("Rows", ["Name", "Value"], [['=HYPERLINK("x")', -5], ["@cmd", 2]])],
    )
    rows = _csv_rows(to_csv(report))
    assert rows[1] == ['\'=HYPERLINK("x")', "-5"] and rows[2][0] == "'@cmd"
    wb = load_workbook(io.BytesIO(to_xlsx(report)))
    assert wb["Rows"]["A2"].value.startswith("'=") and wb["Rows"]["B2"].value == -5
