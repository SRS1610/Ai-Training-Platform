"""Knowledge base: upload → extract → chunk → embed (pgvector) → retrieval; approval; SSRF protection."""

from __future__ import annotations

import io

import httpx
import pytest
from sqlalchemy import func, select

from app.models.ai import KnowledgeChunk
from app.modules.ai import gateway
from app.modules.knowledge import ssrf
from app.workers import drain
from tests.ai_helpers import add_kb_text, bypass_session, make_course
from tests.factories import add_member, signup

DNS_TEXT = """# Networking basics

## DNS

DNS translates hostnames such as intranet.example.com into IP addresses. When the DNS server is unreachable or
returns a wrong record, users cannot open websites by name even though the network connection itself works.
Flushing the resolver cache with ipconfig /flushdns and checking the configured DNS server are first steps.

## DHCP

DHCP leases IP addresses to devices automatically. An exhausted DHCP scope prevents new devices from joining.
"""
PRINTER_TEXT = """# Printer maintenance

Paper jams are cleared by opening tray two and removing crumpled sheets. Toner cartridges are replaced when the
status light blinks amber. Always power off the printer before opening the fuser compartment.
"""


def pdf_bytes() -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    y = 800
    for line in ["Password policy", "Passwords must be at least fourteen characters long.",
                 "Multi-factor authentication is mandatory for remote access.",
                 "Never share passwords by email or chat."]:
        c.drawString(72, y, line)
        y -= 20
    c.showPage()
    c.drawString(72, 800, "Report suspected phishing with the Report Phish button.")
    c.save()
    return buf.getvalue()


def docx_bytes() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Incident escalation procedure", level=1)
    d.add_paragraph("Priority one incidents must be escalated to the duty manager within fifteen minutes.")
    d.add_heading("Communication", level=2)
    d.add_paragraph("Send a status update to affected users every thirty minutes until resolution.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


async def test_upload_pdf_and_docx_ingest_and_retrieve(client):
    admin = await signup(client)
    for name, data, ctype in [
        ("policy.pdf", pdf_bytes(), "application/pdf"),
        ("escalation.docx", docx_bytes(),
         "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ]:
        r = await client.post("/api/v1/knowledge/documents/upload", headers=admin.headers,
                              files={"file": (name, data, ctype)}, data={"category": "policy"})
        assert r.status_code == 202, r.text
        assert r.json()["status"] == "pending" and r.json()["approved"] is False
    await drain()
    r = await client.get("/api/v1/knowledge/documents", headers=admin.headers)
    docs = {d["title"]: d for d in r.json()["items"]}
    assert set(docs) == {"policy", "escalation"}
    for d in docs.values():
        assert d["status"] == "ready" and d["chunk_count"] >= 1 and d["char_count"] > 50, d
    assert docs["policy"]["source_type"] == "pdf" and docs["escalation"]["source_type"] == "docx"

    r = await client.get(f"/api/v1/knowledge/documents/{docs['escalation']['id']}/chunks", headers=admin.headers)
    chunk = r.json()["items"][0]
    assert chunk["embedding_model"] == "offline-hash-embedding" and "duty manager" in chunk["content"]
    assert chunk["heading"].startswith("Incident escalation procedure")

    r = await client.post("/api/v1/knowledge/search", headers=admin.headers,
                          json={"query": "how long must passwords be?", "k": 3})
    assert r.status_code == 200
    hits = r.json()["results"]
    assert hits and hits[0]["title"] == "policy" and "fourteen" in hits[0]["snippet"]
    r = await client.post("/api/v1/knowledge/search", headers=admin.headers,
                          json={"query": "escalate priority one incident to whom", "k": 3})
    assert r.json()["results"][0]["title"] == "escalation"
    # approved_only hides unapproved documents
    r = await client.post("/api/v1/knowledge/search", headers=admin.headers,
                          json={"query": "escalate priority one incident", "approved_only": True})
    assert r.json()["results"] == []


async def test_retrieval_relevance_course_scope_and_approval(client):
    admin = await signup(client)
    course = await make_course(admin.org_id, admin.user_id, publish=False)
    dns = await add_kb_text(client, admin.headers, "Networking guide", DNS_TEXT, course_id=course["course_id"])
    await add_kb_text(client, admin.headers, "Printer guide", PRINTER_TEXT)
    other = await add_kb_text(client, admin.headers, "Unapproved DNS notes", DNS_TEXT, approve=False)
    async with bypass_session() as s:
        hits = await gateway.retrieve(s, organisation_id=admin.org_id, query="Why would DNS cause websites to fail?",
                                      course_id=course["course_id"], k=3)
        assert hits[0].document_id.hex == dns["id"].replace("-", "")
        assert hits[0].course_id == course["course_id"] and hits[0].score > 0.3
        assert all(h.title != "Unapproved DNS notes" for h in hits)
        printer = await gateway.retrieve(s, organisation_id=admin.org_id, query="toner cartridge amber light", k=2)
        assert printer[0].title == "Printer guide"
        allhits = await gateway.retrieve(s, organisation_id=admin.org_id, query="DNS resolver cache", k=5,
                                         approved_only=False)
        assert other["id"] in {str(h.document_id) for h in allhits}
        # keyword fallback still finds exact terms
        kw = await gateway.retrieve(s, organisation_id=admin.org_id, query="flushdns", k=2)
        assert kw and "flushdns" in kw[0].content
        results = await gateway.semantic_search(s, organisation_id=admin.org_id, query="DHCP scope exhausted", k=5)
        assert results[0]["type"] == "knowledge" and results[0]["title"] in ("Networking guide",
                                                                             "Unapproved DNS notes")
        assert results[0]["link"].startswith("/app/instructor/knowledge")
        assert len({r["document_id"] for r in results}) == len(results)  # one hit per document


async def test_lifecycle_approve_reprocess_update_delete_and_permissions(client):
    admin = await signup(client)
    doc = await add_kb_text(client, admin.headers, "Networking guide", DNS_TEXT, approve=False)
    did = doc["id"]
    r = await client.post(f"/api/v1/knowledge/documents/{did}/approve", headers=admin.headers)
    assert r.json()["approved"] and r.json()["approved_by"] == str(admin.user_id)
    r = await client.post(f"/api/v1/knowledge/documents/{did}/unapprove", headers=admin.headers)
    assert not r.json()["approved"]
    course = await make_course(admin.org_id, admin.user_id, publish=False)
    r = await client.patch(f"/api/v1/knowledge/documents/{did}", headers=admin.headers,
                           json={"course_id": str(course["course_id"]), "title": "DNS & DHCP"})
    assert r.json()["course_id"] == str(course["course_id"]) and r.json()["title"] == "DNS & DHCP"
    async with bypass_session() as s:
        cids = set((await s.scalars(select(KnowledgeChunk.course_id).where(
            KnowledgeChunk.document_id == did))).all())
    assert cids == {course["course_id"]}
    r = await client.post(f"/api/v1/knowledge/documents/{did}/reprocess", headers=admin.headers)
    assert r.status_code == 202 and r.json()["status"] == "pending"
    await drain()
    r = await client.get(f"/api/v1/knowledge/documents/{did}", headers=admin.headers)
    assert r.json()["status"] == "ready"

    learner = await add_member(admin.org_id, "learner")
    assert (await client.get("/api/v1/knowledge/documents", headers=learner.headers)).status_code == 403
    assert (await client.post("/api/v1/knowledge/search", headers=learner.headers,
                              json={"query": "dns"})).status_code == 403
    instructor = await add_member(admin.org_id, "instructor")
    assert (await client.get("/api/v1/knowledge/documents", headers=instructor.headers)).status_code == 200

    other_org = await signup(client, org_name="Other Org")
    assert (await client.get(f"/api/v1/knowledge/documents/{did}", headers=other_org.headers)).status_code == 404

    r = await client.delete(f"/api/v1/knowledge/documents/{did}", headers=admin.headers)
    assert r.status_code == 204
    async with bypass_session() as s:
        n = await s.scalar(select(func.count()).select_from(KnowledgeChunk).where(KnowledgeChunk.document_id == did))
    assert n == 0
    assert (await client.get(f"/api/v1/knowledge/documents/{did}", headers=admin.headers)).status_code == 404


async def test_upload_rejects_bad_types_and_failed_extraction_is_captured(client):
    admin = await signup(client)
    r = await client.post("/api/v1/knowledge/documents/upload", headers=admin.headers,
                          files={"file": ("evil.exe", b"MZ....", "application/octet-stream")})
    assert r.status_code == 400
    r = await client.post("/api/v1/knowledge/documents/upload", headers=admin.headers,
                          files={"file": ("image.png", b"\x89PNG\r\n", "image/png")})
    assert r.status_code == 400
    # A structurally broken PDF passes the upload checks but fails extraction → status failed + error
    r = await client.post("/api/v1/knowledge/documents/upload", headers=admin.headers,
                          files={"file": ("broken.pdf", b"%PDF-1.4 garbage", "application/pdf")})
    assert r.status_code == 202
    await drain()
    r = await client.get(f"/api/v1/knowledge/documents/{r.json()['id']}", headers=admin.headers)
    assert r.json()["status"] == "failed" and r.json()["error"]


@pytest.mark.parametrize("url", [
    "http://127.0.0.1/admin",
    "http://169.254.169.254/latest/meta-data/iam/security-credentials/",
    "http://localhost:8000/api",
    "http://10.0.0.5/",
    "http://192.168.1.1/",
    "http://[::1]/",
    "http://[::ffff:127.0.0.1]/",
    "http://0.0.0.0/",
    "http://100.64.0.1/",
    "file:///etc/passwd",
    "ftp://example.com/file.pdf",
    "gopher://example.com/",
    "http://user:pass@example.com/",
    "http://example.com:22/",
    "http://metadata.google.internal/computeMetadata/v1/",
    "http://printer.local/",
])
async def test_ssrf_blocked_urls(client, url):
    admin = await signup(client)
    r = await client.post("/api/v1/knowledge/documents/url", headers=admin.headers, json={"url": url})
    assert r.status_code == 400, (url, r.text)
    assert r.json()["error"]["code"] == "unsafe_url"


async def test_ssrf_dns_resolution_and_redirect_checks(client, monkeypatch):
    admin = await signup(client)
    dns = {"public.example.org": ["93.184.216.34"], "sneaky.example.org": ["10.1.2.3"],
           "mixed.example.org": ["93.184.216.34", "127.0.0.1"], "internal.example.org": ["172.16.0.10"]}

    async def fake_resolve(host, port):
        return dns.get(host, ["93.184.216.34"])

    monkeypatch.setattr(ssrf, "resolve_host", fake_resolve)
    for host in ("sneaky.example.org", "mixed.example.org"):
        r = await client.post("/api/v1/knowledge/documents/url", headers=admin.headers,
                              json={"url": f"https://{host}/kb"})
        assert r.status_code == 400 and r.json()["error"]["code"] == "unsafe_url"

    html = (b"<html><head><title>VPN guide</title></head><body><main><h1>VPN</h1><p>Connect with the corporate "
            b"VPN client before accessing the intranet. Use your MFA token when prompted.</p></main></body></html>")

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.host == "public.example.org" and request.url.path == "/redirect":
            return httpx.Response(302, headers={"location": "http://internal.example.org/secret"})
        if request.url.host == "public.example.org" and request.url.path == "/big":
            return httpx.Response(200, headers={"content-type": "text/plain"}, content=b"a" * (ssrf.MAX_BYTES + 10))
        if request.url.host == "public.example.org" and request.url.path == "/bin":
            return httpx.Response(200, headers={"content-type": "application/x-msdownload"}, content=b"MZ")
        if request.url.host == "internal.example.org":
            raise AssertionError("must never connect to an internal host")
        return httpx.Response(200, headers={"content-type": "text/html; charset=utf-8"}, content=html)

    monkeypatch.setattr(ssrf, "fetch_transport", httpx.MockTransport(handler))
    r = await client.post("/api/v1/knowledge/documents/url", headers=admin.headers,
                          json={"url": "https://public.example.org/vpn"})
    assert r.status_code == 202, r.text
    ok_id = r.json()["id"]
    bad = {}
    for path in ("redirect", "big", "bin"):
        r = await client.post("/api/v1/knowledge/documents/url", headers=admin.headers,
                              json={"url": f"https://public.example.org/{path}"})
        assert r.status_code == 202
        bad[path] = r.json()["id"]
    await drain()
    r = await client.get(f"/api/v1/knowledge/documents/{ok_id}", headers=admin.headers)
    assert r.json()["status"] == "ready" and r.json()["title"] == "VPN guide" and r.json()["source_type"] == "url"
    errors = {}
    for path, did in bad.items():
        d = (await client.get(f"/api/v1/knowledge/documents/{did}", headers=admin.headers)).json()
        assert d["status"] == "failed", d
        errors[path] = d["error"]
    assert "private or internal" in errors["redirect"]
    assert "too large" in errors["big"] and "Unsupported content type" in errors["bin"]


def test_ip_blocklist_unit():
    import ipaddress

    for ip in ("127.0.0.1", "10.0.0.1", "172.16.5.4", "192.168.0.1", "169.254.169.254", "100.64.1.1", "0.0.0.0",
               "::1", "fe80::1", "fc00::1", "::ffff:10.0.0.1", "64:ff9b::a00:1", "224.0.0.1", "255.255.255.255"):
        assert ssrf.ip_is_blocked(ipaddress.ip_address(ip)), ip
    for ip in ("93.184.216.34", "8.8.8.8", "2606:4700:4700::1111"):
        assert not ssrf.ip_is_blocked(ipaddress.ip_address(ip)), ip
